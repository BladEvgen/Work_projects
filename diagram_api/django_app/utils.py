import json
import os
from django.conf import settings
import mysql.connector
import openpyxl
from dotenv import load_dotenv

dotenv_path = os.path.join(os.path.dirname(__file__), ".env")
if os.path.exists(dotenv_path):
    load_dotenv(dotenv_path)


def count_debtors(file):
    wb: openpyxl.Workbook = openpyxl.load_workbook(file)
    result: dict[str, dict] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name != "Лист1":
            sheet: openpyxl.Worksheet = wb[sheet_name]
            data: list = [cell.value for cell in sheet["B"] if cell.value is not None]
            result[transliterate(sheet_name.lower())] = len(data)
    return result


def execute_query(cursor, query, params=None):
    cursor.execute(query, params) if params else cursor.execute(query)
    return cursor.fetchall()


def connect_to_database(database):
    return mysql.connector.connect(
        host=os.getenv("DB_HOST"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"),
        database=database,
    )


def get_tutor_ids(cursor, ignore_vdovtsev=True):
    query = "SELECT DISTINCT TutorID FROM tutor_rating"
    if ignore_vdovtsev:
        query += " WHERE TutorID != 861"
    return [row[0] for row in execute_query(cursor, query)]


def get_fio(cursor, tutor_id):
    result = execute_query(
        cursor,
        "SELECT lastname, firstname, patronymic FROM tutors WHERE TutorID = %s",
        (tutor_id,),
    )
    if result:
        lastname, firstname, patronymic = result[0]
        return (
            f"{lastname} {firstname} {patronymic}"
            if patronymic
            else f"{lastname} {firstname}"
        )
    else:
        return None


def get_comments_count(cursor, tutor_id):
    return execute_query(
        cursor,
        'SELECT COUNT(*) FROM rating_students_vote WHERE TutorID = %s AND json NOT LIKE \'%%"professor_comment":""%%\'',
        (tutor_id,),
    )[0][0]


def get_votes_count(cursor, tutor_id):
    return execute_query(
        cursor, "SELECT SUM(voted) FROM tutor_rating WHERE TutorID = %s", (tutor_id,)
    )[0][0]


def get_rating(cursor, tutor_id):
    result = execute_query(
        cursor, "SELECT rating, voted FROM tutor_rating WHERE TutorID = %s", (tutor_id,)
    )
    if result:
        rating, voted = result[0]
        if rating is None or voted is None or voted == 0:
            return 0.0, "mid", 0
        else:
            rating = float(rating)
            gpa_all = round(rating / (voted * 5), 1)
            grade_class = "good" if gpa_all > 4.0 else "bad" if gpa_all < 3.0 else "mid"
            return gpa_all, grade_class, voted
    else:
        return 0.0, "mid", 0


def get_comments(cursor, tutor_id):
    comments = []
    result = execute_query(
        cursor, "SELECT json FROM rating_students_vote WHERE TutorID = %s", (tutor_id,)
    )
    for row in result:
        try:
            comment_data = json.loads(row[0])
            professor_comment = comment_data.get("professor_comment", "")
            if professor_comment:
                comments.append(professor_comment)
        except json.JSONDecodeError:
            pass
    return comments


def write_to_excel(data, comments_data, comments_enabled=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = [
        "ФИО",
        "Рейтинг",
        "Количество голосов",
        "Количество комментариев",
    ]
    filename = "tutor_raiting.xlsx"

    if comments_enabled:
        headers.append("Комментарии")
        filename = "tutor_rating_comments.xlsx"

    ws.append(headers)

    for tutor_data, comments in zip(data, comments_data):
        fio, gpa_rating, votes_count, comments_count = tutor_data
        row = [
            fio,
            gpa_rating,
            votes_count,
            comments_count,
        ]
        if comments_enabled:
            row.append(comments[0] if comments else "")
        ws.append(row)
        for comment in comments[1:]:
            ws.append(["", "", "", "", comment])

    file_path = os.path.join(settings.STATIC_ROOT, filename)
    wb.save(file_path)


def get_comments_from_db(comments_enabled=True, ignore_vdovtsev=True):
    with connect_to_database("users_base") as users_base_conn:
        with users_base_conn.cursor() as users_base_cursor:
            tutor_ids = get_tutor_ids(users_base_cursor, ignore_vdovtsev)
            data = []
            comments_data = []

            with connect_to_database("platonus") as platonus_conn:
                with platonus_conn.cursor() as platonus_cursor:
                    for tutor_id in tutor_ids:
                        fio = get_fio(platonus_cursor, tutor_id)
                        if fio:
                            comments_count = get_comments_count(
                                users_base_cursor, tutor_id
                            )
                            votes_count = get_votes_count(users_base_cursor, tutor_id)
                            gpa_rating, grade_class, _ = get_rating(
                                users_base_cursor, tutor_id
                            )
                            if comments_enabled:
                                comments = get_comments(users_base_cursor, tutor_id)
                            else:
                                comments = []
                            data.append([fio, gpa_rating, votes_count, comments_count])
                            comments_data.append(comments)

            write_to_excel(data, comments_data, comments_enabled)


def transliterate(name):
    slovar = {
        "а": "a",
        "б": "b",
        "в": "v",
        "г": "g",
        "д": "d",
        "е": "e",
        "ё": "yo",
        "ж": "zh",
        "з": "z",
        "и": "i",
        "й": "y",
        "к": "k",
        "л": "l",
        "м": "m",
        "н": "n",
        "о": "o",
        "п": "p",
        "р": "r",
        "с": "s",
        "т": "t",
        "у": "u",
        "ф": "f",
        "х": "h",
        "ц": "ts",
        "ч": "ch",
        "ш": "sh",
        "щ": "sch",
        "ъ": "",
        "ы": "y",
        "ь": "",
        "э": "e",
        "ю": "yu",
        "я": "ya",
        " ": " ",
        "-": "-",
        ".": ".",
        ",": ",",
        "!": "!",
        "?": "?",
        ":": ":",
    }

    name = name.lower()

    translit = ""
    for letter in name:
        if letter in slovar:
            translit += slovar[letter]
        else:
            translit += letter

    return translit
